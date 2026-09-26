"""Render the canonical evaluation Markdown as a readable Word report."""

from __future__ import annotations

from hashlib import sha256
from pathlib import Path
from zipfile import BadZipFile


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "EVALUATION_REPORT.md"
OUTPUT = ROOT / "docs" / "EVALUATION_REPORT.docx"


def _source_fingerprint() -> str:
    return sha256(SOURCE.read_bytes() + Path(__file__).read_bytes()).hexdigest()


def docx_current() -> bool:
    if not OUTPUT.exists():
        return False
    try:
        from docx import Document

        return Document(OUTPUT).core_properties.identifier == _source_fingerprint()
    except (BadZipFile, OSError, ValueError):
        return False


def _cell_shading(cell, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def _cell_margins(cell) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for side, value in (("top", "70"), ("bottom", "70"), ("left", "85"), ("right", "85")):
        element = OxmlElement(f"w:{side}")
        element.set(qn("w:w"), value)
        element.set(qn("w:type"), "dxa")
        margins.append(element)
    tc_pr.append(margins)


def _set_borders(table) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    table_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = OxmlElement(f"w:{edge}")
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "4")
        item.set(qn("w:color"), "D9D9D9")
        borders.append(item)
    table_pr.append(borders)


def _append_inline(paragraph, token) -> None:
    from docx.shared import Pt

    bold = italic = False
    for child in token.children or []:
        kind = child.type
        if kind == "strong_open":
            bold = True
        elif kind == "strong_close":
            bold = False
        elif kind == "em_open":
            italic = True
        elif kind == "em_close":
            italic = False
        elif kind in {"text", "code_inline", "softbreak", "hardbreak"}:
            value = " " if kind == "softbreak" else "\n" if kind == "hardbreak" else child.content
            run = paragraph.add_run(value)
            run.bold = bold
            run.italic = italic
            if kind == "code_inline":
                run.font.name = "Consolas"
                run.font.size = Pt(8.5)


def _add_table(document, rows: list[list]) -> None:
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.shared import Inches

    columns = max(map(len, rows))
    widths = ([1.22] + [(7.1 - 1.22) / (columns - 1)] * (columns - 1)
              if columns >= 5 else [4.55, 2.55] if columns == 2 else [7.1 / columns] * columns)
    table = document.add_table(rows=0, cols=columns)
    table.autofit = False
    table.style = "Table Grid"
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    _set_borders(table)
    for index, cells in enumerate(rows):
        row = table.add_row()
        if index == 0:
            row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for column, cell in enumerate(row.cells):
            cell.width = Inches(widths[column])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _cell_margins(cell)
            if index == 0:
                _cell_shading(cell, "E9EFF6")
            elif index % 2 == 0:
                _cell_shading(cell, "F7F9FC")
            paragraph = cell.paragraphs[0]
            paragraph.style = document.styles["Table Content"]
            if column > 0:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if column < len(cells):
                _append_inline(paragraph, cells[column])
            for run in paragraph.runs:
                if index == 0:
                    run.bold = True
    document.add_paragraph().paragraph_format.space_after = 0


def build_docx() -> Path:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from markdown_it import MarkdownIt

    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    for name, size, before, after in (
        ("Title", 19, 0, 13),
        ("Heading 1", 13.5, 13, 6),
        ("Heading 2", 11.5, 10, 5),
    ):
        style = document.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    table_style = document.styles.add_style("Table Content", 1)
    table_style.base_style = normal
    table_style.font.size = Pt(8)
    table_style.paragraph_format.space_after = Pt(0)
    table_style.paragraph_format.line_spacing = 1.05
    bullet_style = document.styles["List Bullet"]
    bullet_style.font.name = "Arial"
    bullet_style.font.size = Pt(10)
    bullet_style.paragraph_format.space_after = Pt(3)

    tokens = MarkdownIt("commonmark").enable("table").parse(SOURCE.read_text(encoding="utf-8"))
    list_depth = 0
    index = 0
    while index < len(tokens):
        token = tokens[index]
        kind = token.type
        if kind == "heading_open":
            level = min(int(token.tag[1:]), 3)
            style = "Title" if level == 1 else "Heading 1" if level == 2 else "Heading 2"
            paragraph = document.add_paragraph(style=style)
            _append_inline(paragraph, tokens[index + 1])
            index += 2
        elif kind == "paragraph_open":
            paragraph = document.add_paragraph(style="List Bullet" if list_depth else "Normal")
            _append_inline(paragraph, tokens[index + 1])
            index += 2
        elif kind == "bullet_list_open":
            list_depth += 1
        elif kind == "bullet_list_close":
            list_depth -= 1
        elif kind == "table_open":
            rows: list[list] = []
            row: list = []
            index += 1
            while index < len(tokens) and tokens[index].type != "table_close":
                item = tokens[index]
                if item.type == "tr_open":
                    row = []
                elif item.type == "inline":
                    row.append(item)
                elif item.type == "tr_close":
                    rows.append(row)
                index += 1
            _add_table(document, rows)
        index += 1

    document.core_properties.title = "PII Redaction Evaluation Report"
    document.core_properties.author = ""
    document.core_properties.identifier = _source_fingerprint()
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    return OUTPUT


def main() -> int:
    print(f"Created {build_docx()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
