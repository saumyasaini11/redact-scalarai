from pathlib import Path

from docx import Document

from pii_redactor.docx_io import DocxPackage, extracted_text
from pii_redactor.models import PIIRecord, PIIType
from pii_redactor.qa import scan_original_values


def test_cross_run_replacement():
    work = Path(".tmp")
    work.mkdir(exist_ok=True)
    source = work / "test-cross-run-source.docx"
    output = work / "test-cross-run-output.docx"
    doc = Document()
    paragraph = doc.add_paragraph()
    paragraph.add_run("john.")
    paragraph.add_run("smith@example.com")
    doc.save(source)

    try:
        package = DocxPackage(source)
        blocks = package.extract_blocks()
        target = blocks[0]
        record = PIIRecord(
            record_id="email", pii_type=PIIType.EMAIL,
            original_text=target.text, normalized_text=target.text,
            source_kind="native_text", document_part=target.part_name,
            block_id=target.block_id, start_offset=0, end_offset=len(target.text),
            replacement_value="user@example.test",
        )
        package.apply_records([record])
        package.save(output)
        assert "user@example.test" in extracted_text(output)
        assert "john.smith@example.com" not in extracted_text(output)
    finally:
        source.unlink(missing_ok=True)
        output.unlink(missing_ok=True)


def test_fallback_does_not_redact_inside_generated_replacement():
    work = Path(".tmp")
    work.mkdir(exist_ok=True)
    source = work / "test-protected-replacement-source.docx"
    output = work / "test-protected-replacement-output.docx"
    doc = Document()
    doc.add_paragraph("Acme Private Limited")
    doc.add_paragraph("Private Limited")
    doc.save(source)

    try:
        package = DocxPackage(source)
        first, second = package.extract_blocks()
        records = [
            PIIRecord(
                record_id="full-company", pii_type=PIIType.COMPANY,
                original_text=first.text, normalized_text=first.text.casefold(),
                source_kind="native_text", document_part=first.part_name,
                block_id=first.block_id, start_offset=0, end_offset=len(first.text),
                replacement_value="Oakbridge Private Limited",
            ),
            PIIRecord(
                record_id="suffix", pii_type=PIIType.COMPANY,
                original_text=second.text, normalized_text=second.text.casefold(),
                source_kind="native_text", document_part=second.part_name,
                block_id=second.block_id, start_offset=0, end_offset=len(second.text),
                replacement_value="[COMPANY]",
            ),
        ]
        package.apply_records(records)
        package.save(output)
        text = extracted_text(output)
        assert "Oakbridge Private Limited" in text
        assert "Oakbridge [COMPANY]" not in text
    finally:
        source.unlink(missing_ok=True)
        output.unlink(missing_ok=True)


def test_residual_scan_is_limited_to_the_original_block():
    work = Path(".tmp")
    work.mkdir(exist_ok=True)
    output = work / "test-block-scoped-residual.docx"
    doc = Document()
    doc.add_paragraph("[COMPANY]")
    doc.add_paragraph("General Company information")
    doc.save(output)

    try:
        block = DocxPackage(output).extract_blocks()[0]
        record = PIIRecord(
            record_id="company", pii_type=PIIType.COMPANY,
            original_text="Company", normalized_text="company",
            source_kind="native_text", document_part=block.part_name,
            block_id=block.block_id, start_offset=0, end_offset=7,
            replacement_value="[COMPANY]",
        )
        assert scan_original_values(output, [record]) == []
    finally:
        output.unlink(missing_ok=True)


def test_fallback_replaces_cross_run_repeats_without_touching_substrings():
    work = Path(".tmp")
    work.mkdir(exist_ok=True)
    source = work / "test-fallback-boundaries-source.docx"
    output = work / "test-fallback-boundaries-output.docx"
    doc = Document()
    first = doc.add_paragraph()
    first.add_run("Alpha ")
    first.add_run("Company")
    second = doc.add_paragraph()
    second.add_run("Alpha ")
    second.add_run("Company")
    doc.add_paragraph("Accompanying text remains intact")
    doc.save(source)

    try:
        package = DocxPackage(source)
        block = package.extract_blocks()[0]
        record = PIIRecord(
            record_id="company", pii_type=PIIType.COMPANY,
            original_text="Alpha Company", normalized_text="alpha company",
            source_kind="native_text", document_part=block.part_name,
            block_id=block.block_id, start_offset=0, end_offset=len(block.text),
            replacement_value="[COMPANY]",
        )
        package.apply_records([record])
        package.save(output)
        text = extracted_text(output)
        assert text.count("[COMPANY]") == 2
        assert "Accompanying text remains intact" in text
    finally:
        source.unlink(missing_ok=True)
        output.unlink(missing_ok=True)


def test_tables_headers_footers_and_run_formatting_are_preserved():
    work = Path(".tmp")
    work.mkdir(exist_ok=True)
    source = work / "test-structure-source.docx"
    output = work / "test-structure-output.docx"
    doc = Document()
    paragraph = doc.add_paragraph()
    first_run = paragraph.add_run("alpha@")
    first_run.bold = True
    paragraph.add_run("example.test")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "table@example.test"
    doc.sections[0].header.paragraphs[0].text = "header@example.test"
    doc.sections[0].footer.paragraphs[0].text = "footer@example.test"
    doc.save(source)

    try:
        package = DocxPackage(source)
        blocks = package.extract_blocks()
        records = []
        for item in blocks:
            if "@example.test" not in item.text:
                continue
            records.append(PIIRecord(
                record_id=item.block_id,
                pii_type=PIIType.EMAIL,
                original_text=item.text,
                normalized_text=item.text.casefold(),
                source_kind="native_text",
                document_part=item.part_name,
                block_id=item.block_id,
                start_offset=0,
                end_offset=len(item.text),
                replacement_value=f"safe-{len(records)}@example.test",
            ))
        package.apply_records(records)
        package.save(output)

        result = Document(output)
        assert result.paragraphs[0].runs[0].bold is True
        assert result.tables[0].cell(0, 0).text.startswith("safe-")
        assert result.sections[0].header.paragraphs[0].text.startswith("safe-")
        assert result.sections[0].footer.paragraphs[0].text.startswith("safe-")
        assert not scan_original_values(output, records)
    finally:
        source.unlink(missing_ok=True)
        output.unlink(missing_ok=True)
